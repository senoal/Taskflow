from __future__ import annotations

import re
import sqlite3
from io import BytesIO
from contextlib import closing
from datetime import datetime
from pathlib import Path
from urllib.parse import urlparse

from flask import Flask, abort, jsonify, render_template, request, send_file, send_from_directory
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from werkzeug.utils import secure_filename

app = Flask(__name__)
DB_PATH = Path(__file__).with_name("pythones_todo.db")
UPLOAD_DIR = Path(__file__).with_name("uploads")
MAX_FILE_SIZE = 50 * 1024 * 1024
APP_VERSION = "v1.6.0"
LAST_UPDATED = "2026-09-30"
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE


def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def normalize_link(value):
    link = str(value or "").strip()
    if not link:
        return ""
    if len(link) > 2048:
        raise ValueError("Link maksimal 2048 karakter.")
    if any(character.isspace() for character in link):
        raise ValueError("Link tidak boleh mengandung spasi.")
    if re.match(r"^[a-z][a-z0-9+.-]*:", link, re.IGNORECASE) and not re.match(r"^https?://", link, re.IGNORECASE):
        raise ValueError("Link harus berupa alamat web HTTP atau HTTPS yang valid.")
    if not re.match(r"^https?://", link, re.IGNORECASE):
        link = f"https://{link}"
    parsed = urlparse(link)
    if parsed.scheme not in {"http", "https"} or not parsed.netloc:
        raise ValueError("Link harus berupa alamat web HTTP atau HTTPS yang valid.")
    return link


def init_db():
    UPLOAD_DIR.mkdir(exist_ok=True)
    with closing(get_db()) as conn:
        conn.execute(
            """CREATE TABLE IF NOT EXISTS tasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title TEXT NOT NULL,
                description TEXT DEFAULT '',
                link TEXT DEFAULT '',
                category TEXT DEFAULT 'Personal',
                due_date TEXT,
                priority TEXT NOT NULL DEFAULT 'medium',
                completed INTEGER NOT NULL DEFAULT 0,
                canceled INTEGER NOT NULL DEFAULT 0,
                created_at TEXT NOT NULL
            )"""
        )
        columns = {row[1] for row in conn.execute("PRAGMA table_info(tasks)").fetchall()}
        if "canceled" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN canceled INTEGER NOT NULL DEFAULT 0")
        if "link" not in columns:
            conn.execute("ALTER TABLE tasks ADD COLUMN link TEXT DEFAULT ''")
        conn.execute(
            """CREATE TABLE IF NOT EXISTS attachments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL UNIQUE,
                size INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL,
                FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE
            )"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS subtasks (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                task_id INTEGER NOT NULL,
                name TEXT NOT NULL,
                due_date TEXT,
                notes TEXT DEFAULT '',
                created_at TEXT NOT NULL,
                FOREIGN KEY(task_id) REFERENCES tasks(id) ON DELETE CASCADE
            )"""
        )
        conn.execute(
            """CREATE TABLE IF NOT EXISTS subtask_attachments (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                subtask_id INTEGER NOT NULL,
                original_name TEXT NOT NULL,
                stored_name TEXT NOT NULL UNIQUE,
                size INTEGER NOT NULL,
                uploaded_at TEXT NOT NULL,
                FOREIGN KEY(subtask_id) REFERENCES subtasks(id) ON DELETE CASCADE
            )"""
        )
        conn.commit()


@app.get("/")
def index():
    return render_template("index.html", app_version=APP_VERSION, last_updated=LAST_UPDATED)


@app.get("/api/tasks")
def get_tasks():
    status = request.args.get("status", "all")
    paginate = "page" in request.args
    search = request.args.get("search", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    try:
        page = max(1, int(request.args.get("page", 1)))
    except ValueError:
        page = 1
    per_page = 5
    with closing(get_db()) as conn:
        query = "SELECT * FROM tasks"
        params = []
        conditions = []
        if status == "active":
            conditions.append("completed = 0 AND canceled = 0")
        elif status == "completed":
            conditions.append("completed = 1 AND canceled = 0")
        elif status == "canceled":
            conditions.append("canceled = 1")
        elif status == "deadline":
            conditions.append("completed = 0 AND canceled = 0 AND due_date = date('now', 'localtime')")
        if search:
            conditions.append("(title LIKE ? OR description LIKE ? OR category LIKE ? OR link LIKE ?)")
            term = f"%{search}%"
            params.extend([term, term, term, term])
        if date_from:
            conditions.append("due_date >= ?")
            params.append(date_from)
        if date_to:
            conditions.append("due_date <= ?")
            params.append(date_to)
        if conditions:
            query += " WHERE " + " AND ".join(conditions)
        count_query = "SELECT COUNT(*) FROM tasks" + (" WHERE " + " AND ".join(conditions) if conditions else "")
        total = conn.execute(count_query, params).fetchone()[0]
        total_pages = max(1, (total + per_page - 1) // per_page)
        page = min(page, total_pages)
        query += " ORDER BY created_at DESC, id DESC LIMIT ? OFFSET ?"
        params.extend([per_page, (page - 1) * per_page])
        rows = conn.execute(query, params).fetchall()
    payload = {"items": [dict(row) for row in rows], "page": page, "per_page": per_page, "total": total, "total_pages": total_pages}
    return jsonify(payload if paginate else payload["items"])


@app.get("/api/tasks/export")
def export_tasks():
    status = request.args.get("status", "all")
    search = request.args.get("search", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    conditions, params = [], []
    if status == "active": conditions.append("completed = 0 AND canceled = 0")
    elif status == "completed": conditions.append("completed = 1 AND canceled = 0")
    elif status == "canceled": conditions.append("canceled = 1")
    elif status == "deadline": conditions.append("completed = 0 AND canceled = 0 AND due_date = date('now', 'localtime')")
    if search:
        conditions.append("(title LIKE ? OR description LIKE ? OR category LIKE ? OR link LIKE ?)")
        term = f"%{search}%"; params.extend([term, term, term, term])
    if date_from: conditions.append("due_date >= ?"); params.append(date_from)
    if date_to: conditions.append("due_date <= ?"); params.append(date_to)
    query = "SELECT * FROM tasks" + (" WHERE " + " AND ".join(conditions) if conditions else "") + " ORDER BY created_at DESC, id DESC LIMIT 701"
    with closing(get_db()) as conn:
        tasks = conn.execute(query, params).fetchall()
    if len(tasks) > 700:
        return jsonify({"error": "Ekspor dibatasi maksimal 700 proyek. Persempit filter lalu coba lagi."}), 422
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Proyek"
    headers = ["ID", "Nama Proyek", "Status", "Bidang / Kategori", "Prioritas", "Target Penyelesaian", "Link", "Deskripsi", "Dibuat Pada"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for task in tasks:
        status_label = "Dihentikan" if task["canceled"] else "Selesai" if task["completed"] else "Berjalan"
        sheet.append([task["id"], task["title"], status_label, task["category"], task["priority"], task["due_date"] or "", task["link"] or "", task["description"] or "", task["created_at"]])
    widths = [10, 32, 14, 18, 12, 14, 40, 45, 22]
    for index, width in enumerate(widths, 1): sheet.column_dimensions[chr(64 + index)].width = width
    sheet.freeze_panes = "A2"; sheet.auto_filter.ref = sheet.dimensions
    output = BytesIO(); workbook.save(output); output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"portofolio_proyek_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/tasks/<int:task_id>")
def get_task(task_id):
    with closing(get_db()) as conn:
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
        attachments = conn.execute("SELECT id, original_name, size, uploaded_at FROM attachments WHERE task_id = ? ORDER BY id DESC", (task_id,)).fetchall()
        subtasks = conn.execute("SELECT * FROM subtasks WHERE task_id = ? ORDER BY due_date IS NULL, due_date ASC, id ASC", (task_id,)).fetchall()
    if not task:
        return jsonify({"error": "Proyek tidak ditemukan."}), 404
    data = dict(task)
    data["attachments"] = [dict(row) for row in attachments]
    data["subtasks"] = [dict(row) for row in subtasks]
    return jsonify(data)


@app.get("/api/tasks/<int:task_id>/subtasks/<int:subtask_id>")
def get_subtask(task_id, subtask_id):
    with closing(get_db()) as conn:
        subtask = conn.execute("SELECT * FROM subtasks WHERE id = ? AND task_id = ?", (subtask_id, task_id)).fetchone()
        attachments = conn.execute("SELECT id, original_name, size, uploaded_at FROM subtask_attachments WHERE subtask_id = ? ORDER BY id DESC", (subtask_id,)).fetchall()
    if not subtask:
        return jsonify({"error": "Tahapan tidak ditemukan."}), 404
    data = dict(subtask)
    data["attachments"] = [dict(row) for row in attachments]
    return jsonify(data)


@app.post("/api/tasks/<int:task_id>/subtasks")
def create_subtask(task_id):
    data = request.get_json() or {}
    name = str(data.get("name", "")).strip()
    if not name:
        return jsonify({"error": "Nama tahapan wajib diisi."}), 400
    with closing(get_db()) as conn:
        if not conn.execute("SELECT 1 FROM tasks WHERE id = ?", (task_id,)).fetchone():
            return jsonify({"error": "Proyek tidak ditemukan."}), 404
        cursor = conn.execute("INSERT INTO subtasks (task_id, name, due_date, notes, created_at) VALUES (?, ?, ?, ?, ?)", (task_id, name, data.get("due_date") or None, str(data.get("notes", "")).strip(), datetime.now().isoformat(timespec="seconds")))
        conn.commit()
        subtask = conn.execute("SELECT * FROM subtasks WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return jsonify(dict(subtask)), 201


@app.patch("/api/tasks/<int:task_id>/subtasks/<int:subtask_id>")
def update_subtask(task_id, subtask_id):
    data = request.get_json() or {}
    updates = {key: value for key, value in data.items() if key in {"name", "due_date", "notes"}}
    if not updates:
        return jsonify({"error": "Tidak ada perubahan."}), 400
    if "name" in updates and not str(updates["name"]).strip():
        return jsonify({"error": "Nama tahapan wajib diisi."}), 400
    if "due_date" in updates:
        updates["due_date"] = updates["due_date"] or None
    clause = ", ".join(f"{key} = ?" for key in updates)
    with closing(get_db()) as conn:
        result = conn.execute(f"UPDATE subtasks SET {clause} WHERE id = ? AND task_id = ?", [*updates.values(), subtask_id, task_id])
        conn.commit()
        subtask = conn.execute("SELECT * FROM subtasks WHERE id = ?", (subtask_id,)).fetchone()
    if not result.rowcount:
        return jsonify({"error": "Tahapan tidak ditemukan."}), 404
    return jsonify(dict(subtask))


@app.delete("/api/tasks/<int:task_id>/subtasks/<int:subtask_id>")
def delete_subtask(task_id, subtask_id):
    with closing(get_db()) as conn:
        files = conn.execute("SELECT stored_name FROM subtask_attachments WHERE subtask_id = ?", (subtask_id,)).fetchall()
        result = conn.execute("DELETE FROM subtasks WHERE id = ? AND task_id = ?", (subtask_id, task_id))
        conn.execute("DELETE FROM subtask_attachments WHERE subtask_id = ?", (subtask_id,))
        conn.commit()
    for file in files:
        (UPLOAD_DIR / file["stored_name"]).unlink(missing_ok=True)
    return ("", 204) if result.rowcount else (jsonify({"error": "Tahapan tidak ditemukan."}), 404)


@app.post("/api/tasks/<int:task_id>/subtasks/<int:subtask_id>/attachments")
def upload_subtask_attachment(task_id, subtask_id):
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Pilih file untuk diunggah."}), 400
    original_name = secure_filename(file.filename)
    if not original_name:
        return jsonify({"error": "Nama file tidak valid."}), 400
    file.stream.seek(0, 2); size = file.stream.tell(); file.stream.seek(0)
    if size > MAX_FILE_SIZE:
        return jsonify({"error": "Ukuran file maksimal 50 MB."}), 413
    with closing(get_db()) as conn:
        if not conn.execute("SELECT 1 FROM subtasks WHERE id = ? AND task_id = ?", (subtask_id, task_id)).fetchone():
            return jsonify({"error": "Tahapan tidak ditemukan."}), 404
        stored_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{original_name}"
        file.save(UPLOAD_DIR / stored_name)
        cursor = conn.execute("INSERT INTO subtask_attachments (subtask_id, original_name, stored_name, size, uploaded_at) VALUES (?, ?, ?, ?, ?)", (subtask_id, original_name, stored_name, size, datetime.now().isoformat(timespec="seconds")))
        conn.commit()
    return jsonify({"id": cursor.lastrowid, "original_name": original_name, "size": size}), 201


@app.delete("/api/subtask-attachments/<int:attachment_id>")
def delete_subtask_attachment(attachment_id):
    with closing(get_db()) as conn:
        attachment = conn.execute("SELECT stored_name FROM subtask_attachments WHERE id = ?", (attachment_id,)).fetchone()
        if not attachment:
            return jsonify({"error": "Lampiran tidak ditemukan."}), 404
        conn.execute("DELETE FROM subtask_attachments WHERE id = ?", (attachment_id,))
        conn.commit()
    (UPLOAD_DIR / attachment["stored_name"]).unlink(missing_ok=True)
    return "", 204


@app.get("/api/subtask-attachments/<int:attachment_id>/download")
def download_subtask_attachment(attachment_id):
    with closing(get_db()) as conn:
        attachment = conn.execute("SELECT original_name, stored_name FROM subtask_attachments WHERE id = ?", (attachment_id,)).fetchone()
    if not attachment:
        abort(404)
    return send_from_directory(UPLOAD_DIR, attachment["stored_name"], as_attachment=True, download_name=attachment["original_name"])


@app.post("/api/tasks/<int:task_id>/attachments")
def upload_attachment(task_id):
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify({"error": "Pilih file untuk diunggah."}), 400
    original_name = secure_filename(file.filename)
    if not original_name:
        return jsonify({"error": "Nama file tidak valid."}), 400
    with closing(get_db()) as conn:
        if not conn.execute("SELECT 1 FROM tasks WHERE id = ?", (task_id,)).fetchone():
            return jsonify({"error": "Proyek tidak ditemukan."}), 404
        file.stream.seek(0, 2)
        size = file.stream.tell()
        file.stream.seek(0)
        if size > MAX_FILE_SIZE:
            return jsonify({"error": "Ukuran file maksimal 50 MB."}), 413
        stored_name = f"{datetime.now().strftime('%Y%m%d%H%M%S%f')}_{original_name}"
        file.save(UPLOAD_DIR / stored_name)
        cursor = conn.execute("INSERT INTO attachments (task_id, original_name, stored_name, size, uploaded_at) VALUES (?, ?, ?, ?, ?)", (task_id, original_name, stored_name, size, datetime.now().isoformat(timespec="seconds")))
        conn.commit()
    return jsonify({"id": cursor.lastrowid, "original_name": original_name, "size": size}), 201


@app.get("/api/attachments/<int:attachment_id>/download")
def download_attachment(attachment_id):
    with closing(get_db()) as conn:
        attachment = conn.execute("SELECT original_name, stored_name FROM attachments WHERE id = ?", (attachment_id,)).fetchone()
    if not attachment:
        abort(404)
    return send_from_directory(UPLOAD_DIR, attachment["stored_name"], as_attachment=True, download_name=attachment["original_name"])


@app.delete("/api/attachments/<int:attachment_id>")
def delete_attachment(attachment_id):
    with closing(get_db()) as conn:
        attachment = conn.execute("SELECT stored_name FROM attachments WHERE id = ?", (attachment_id,)).fetchone()
        if not attachment:
            return jsonify({"error": "Lampiran tidak ditemukan."}), 404
        conn.execute("DELETE FROM attachments WHERE id = ?", (attachment_id,))
        conn.commit()
    (UPLOAD_DIR / attachment["stored_name"]).unlink(missing_ok=True)
    return "", 204


@app.post("/api/tasks")
def create_task():
    data = request.get_json() or {}
    title = str(data.get("title", "")).strip()
    if not title:
        return jsonify({"error": "Nama proyek wajib diisi."}), 400
    priority = data.get("priority", "medium")
    if priority not in {"low", "medium", "high"}:
        priority = "medium"
    try:
        link = normalize_link(data.get("link", ""))
    except ValueError as error:
        return jsonify({"error": str(error)}), 400
    with closing(get_db()) as conn:
        cursor = conn.execute(
            "INSERT INTO tasks (title, description, link, category, due_date, priority, created_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
            (title, str(data.get("description", "")).strip(), link, str(data.get("category", "Personal")).strip() or "Personal", data.get("due_date") or None, priority, datetime.now().isoformat(timespec="seconds")),
        )
        conn.commit()
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (cursor.lastrowid,)).fetchone()
    return jsonify(dict(task)), 201


@app.patch("/api/tasks/<int:task_id>")
def update_task(task_id):
    data = request.get_json() or {}
    allowed = {"title", "description", "link", "category", "due_date", "priority", "completed", "canceled"}
    updates = {key: value for key, value in data.items() if key in allowed}
    if not updates:
        return jsonify({"error": "Tidak ada perubahan."}), 400
    if "title" in updates and not str(updates["title"]).strip():
        return jsonify({"error": "Nama proyek wajib diisi."}), 400
    if "completed" in updates:
        updates["completed"] = int(bool(updates["completed"]))
        if updates["completed"]:
            updates["canceled"] = 0
    if "canceled" in updates:
        updates["canceled"] = int(bool(updates["canceled"]))
        if updates["canceled"]:
            updates["completed"] = 0
    if "due_date" in updates:
        updates["due_date"] = updates["due_date"] or None
    if "link" in updates:
        try:
            updates["link"] = normalize_link(updates["link"])
        except ValueError as error:
            return jsonify({"error": str(error)}), 400
    clause = ", ".join(f"{key} = ?" for key in updates)
    with closing(get_db()) as conn:
        conn.execute(f"UPDATE tasks SET {clause} WHERE id = ?", [*updates.values(), task_id])
        conn.commit()
        task = conn.execute("SELECT * FROM tasks WHERE id = ?", (task_id,)).fetchone()
    if not task:
        return jsonify({"error": "Proyek tidak ditemukan."}), 404
    return jsonify(dict(task))


@app.delete("/api/tasks/<int:task_id>")
def delete_task(task_id):
    with closing(get_db()) as conn:
        attachments = conn.execute("SELECT stored_name FROM attachments WHERE task_id = ?", (task_id,)).fetchall()
        subtask_attachments = conn.execute("SELECT stored_name FROM subtask_attachments WHERE subtask_id IN (SELECT id FROM subtasks WHERE task_id = ?)", (task_id,)).fetchall()
        result = conn.execute("DELETE FROM tasks WHERE id = ?", (task_id,))
        conn.execute("DELETE FROM attachments WHERE task_id = ?", (task_id,))
        conn.execute("DELETE FROM subtask_attachments WHERE subtask_id IN (SELECT id FROM subtasks WHERE task_id = ?)", (task_id,))
        conn.execute("DELETE FROM subtasks WHERE task_id = ?", (task_id,))
        conn.commit()
    for attachment in attachments:
        (UPLOAD_DIR / attachment["stored_name"]).unlink(missing_ok=True)
    for attachment in subtask_attachments:
        (UPLOAD_DIR / attachment["stored_name"]).unlink(missing_ok=True)
    return ("", 204) if result.rowcount else (jsonify({"error": "Proyek tidak ditemukan."}), 404)


@app.errorhandler(413)
def file_too_large(_error):
    return jsonify({"error": "Ukuran file maksimal 50 MB."}), 413


@app.get("/api/stats")
def stats():
    search = request.args.get("search", "").strip()
    date_from = request.args.get("date_from", "").strip()
    date_to = request.args.get("date_to", "").strip()
    conditions, params = [], []
    if search:
        conditions.append("(title LIKE ? OR description LIKE ? OR category LIKE ? OR link LIKE ?)")
        term = f"%{search}%"
        params.extend([term, term, term, term])
    if date_from:
        conditions.append("due_date >= ?")
        params.append(date_from)
    if date_to:
        conditions.append("due_date <= ?")
        params.append(date_to)
    where = (" WHERE " + " AND ".join(conditions)) if conditions else ""
    with closing(get_db()) as conn:
        total = conn.execute("SELECT COUNT(*) FROM tasks" + where, params).fetchone()[0]
        completed_where = where + (" AND " if where else " WHERE ") + "completed = 1"
        today_where = where + (" AND " if where else " WHERE ") + "completed = 0 AND due_date = date('now', 'localtime')"
        completed = conn.execute("SELECT COUNT(*) FROM tasks" + completed_where + " AND canceled = 0", params).fetchone()[0]
        due_today = conn.execute("SELECT COUNT(*) FROM tasks" + today_where + " AND canceled = 0", params).fetchone()[0]
        canceled = conn.execute("SELECT COUNT(*) FROM tasks" + where + ((" AND " if where else " WHERE ") + "canceled = 1"), params).fetchone()[0]
    return jsonify({"total": total, "completed": completed, "active": total - completed - canceled, "canceled": canceled, "due_today": due_today})


if __name__ == "__main__":
    init_db()
    app.run(debug=True)
